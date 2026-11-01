import json
from datetime import datetime
from pathlib import Path
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from ..config import EXPORT_DIR
from ..models import Job, Company, Evidence, ManualReview, AutomationJob, NationalRoofingJob

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Calibri", size=10)
LINK_FONT = Font(name="Calibri", size=10, color="2563EB", underline="single")
THIN_BORDER = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

def style_header_row(ws, col_count: int):
    ws.freeze_panes = "A2"
    for col_idx in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions

def auto_fit_columns(ws, min_width=12, max_width=45):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if cell.hyperlink:
                val = str(cell.hyperlink.target or val)
            lines = val.split("\n")
            for line in lines:
                if len(line) > max_len:
                    max_len = len(line)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_width), max_width)

def export_job_to_excel(job: Job, companies: list[Company], evidences: list[Evidence], reviews: list[ManualReview]) -> Path:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # ==========================================
    # SHEET 1: Enriched Leads
    # ==========================================
    ws_leads = wb.create_sheet(title="Enriched Leads")
    
    headers_sheet1 = [
        # Original Columns (Preserved 1st)
        "id", "name", "trades", "city", "state", "phone", "webssite url", "claimed",
        # Business Details
        "Official Business Name", "Business Category", "Business Description", "Full Business Address", 
        "ZIP Code", "Country", "Year Established", "Company Size", "Official Website", "Website Status",
        # Agency Intelligence (Tex Web Solutions)
        "Tech Stack / CMS", "SSL Secured", "Mobile Responsive", "Has Analytics", "SEO Health Score", "Redesign Opportunity",
        # Mandatory Decision-Maker Fields (1-4)
        "Owner Name", "Owner Job Title", "CEO Name", "CEO Job Title", 
        "Founder Name", "Founder Job Title", "Co-Founder Name", "Co-Founder Job Title", 
        "Primary Decision Maker", "Primary Decision Maker Title",
        # Mandatory Contact Fields (5-8)
        "Business Email", "Additional Business Emails", "Business Phone", "Additional Business Phone", 
        "Personal Professional Email", "Public Professional Phone", "Phone Line Type", "Email Verification Status", "Mailbox Provider",
        # Mandatory Social Fields (9-12)
        "Personal LinkedIn", "Business LinkedIn", "Business Instagram", "Business Facebook",
        # Research Information
        "Owner Verification Status", "CEO Verification Status", "Founder Verification Status", "Co-Founder Verification Status",
        "Business Email Status", "Business Phone Status", "Personal Email Status", "Personal Phone Status", 
        "Personal LinkedIn Status", "Business LinkedIn Status", "Instagram Status", "Facebook Status", 
        "Overall Confidence Score", "Overall Verification Status", "Missing Fields", "Last Checked", "Research Notes",
        # AI Cold Outreach Ammunition
        "Personalized Outreach Pitch"
    ]
    ws_leads.append(headers_sheet1)

    for comp in companies:
        add_emails = ", ".join(json.loads(comp.additional_business_emails or "[]"))
        add_phones = ", ".join(json.loads(comp.additional_business_phones or "[]"))
        missing = ", ".join(json.loads(comp.missing_fields or "[]"))
        last_chk = comp.last_checked.strftime("%Y-%m-%d %H:%M:%S") if comp.last_checked else ""

        row = [
            comp.original_id, comp.original_name, comp.original_trades, comp.original_city,
            comp.original_state, comp.original_phone, comp.original_website, comp.original_claimed,
            comp.official_name or comp.original_name, comp.business_category or comp.trades,
            comp.business_description, comp.address, comp.zip_code, comp.country,
            comp.year_established, comp.company_size, comp.official_website, comp.website_status,
            comp.tech_stack_cms or "Custom / Unknown",
            "Yes" if comp.has_ssl else "No",
            "Yes" if comp.is_mobile_responsive else "No",
            "Yes" if comp.has_analytics else "No",
            comp.seo_audit_score or 0,
            comp.redesign_opportunity or "Medium",
            comp.owner_name, comp.owner_title, comp.ceo_name, comp.ceo_title,
            comp.founder_name, comp.founder_title, comp.cofounder_name, comp.cofounder_title,
            comp.primary_decision_maker, comp.primary_decision_maker_title,
            comp.business_email, add_emails, comp.business_phone, add_phones,
            comp.personal_email, comp.personal_phone, comp.phone_line_type or "Landline / Office",
            comp.email_verification_status, comp.email_mailbox_provider or "Custom",
            comp.personal_linkedin, comp.business_linkedin, comp.business_instagram, comp.business_facebook,
            comp.owner_status, comp.ceo_status, comp.founder_status, comp.cofounder_status,
            comp.business_email_status, comp.business_phone_status, comp.personal_email_status, comp.personal_phone_status,
            comp.personal_linkedin_status, comp.business_linkedin_status, comp.business_instagram_status, comp.business_facebook_status,
            comp.overall_confidence_score, comp.overall_verification_status, missing, last_chk, comp.research_notes,
            comp.ai_personalized_pitch or ""
        ]
        ws_leads.append(row)
        curr_row = ws_leads.max_row
        
        # Apply formatting and links
        for c_idx in range(1, len(row) + 1):
            cell = ws_leads.cell(row=curr_row, column=c_idx)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            val = str(cell.value or "")
            if val.startswith("http://") or val.startswith("https://"):
                cell.hyperlink = val
                cell.font = LINK_FONT

    style_header_row(ws_leads, len(headers_sheet1))
    auto_fit_columns(ws_leads)

    # ==========================================
    # SHEET 2: Research Evidence
    # ==========================================
    ws_evidence = wb.create_sheet(title="Research Evidence")
    headers_sheet2 = [
        "Business ID", "Business Name", "Field Name", "Discovered Value", 
        "Source Name", "Source URL", "Discovery Date", "Last Checked", 
        "Verification Status", "Confidence Score", "Supporting Evidence"
    ]
    ws_evidence.append(headers_sheet2)

    comp_lookup = {c.id: c for c in companies}
    for ev in evidences:
        comp = comp_lookup.get(ev.company_id)
        b_id = comp.original_id if comp else ""
        b_name = comp.original_name if comp else ""
        disc_date = ev.discovery_date.strftime("%Y-%m-%d %H:%M:%S") if ev.discovery_date else ""
        last_date = ev.last_checked.strftime("%Y-%m-%d %H:%M:%S") if ev.last_checked else ""

        ws_evidence.append([
            b_id, b_name, ev.field_name, ev.discovered_value,
            ev.source_name, ev.source_url, disc_date, last_date,
            ev.verification_status, ev.confidence_score, ev.notes
        ])
        curr_row = ws_evidence.max_row
        for c_idx in range(1, len(headers_sheet2) + 1):
            cell = ws_evidence.cell(row=curr_row, column=c_idx)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if c_idx == 6 and str(cell.value or "").startswith("http"):
                cell.hyperlink = str(cell.value)
                cell.font = LINK_FONT

    style_header_row(ws_evidence, len(headers_sheet2))
    auto_fit_columns(ws_evidence)

    # ==========================================
    # SHEET 3: Manual Review
    # ==========================================
    ws_review = wb.create_sheet(title="Manual Review")
    headers_sheet3 = [
        "Business ID", "Business Name", "Field", "Conflicting Values", 
        "Source URLs", "Reason for Review", "Suggested Action", "Review Status"
    ]
    ws_review.append(headers_sheet3)

    for rev in reviews:
        comp = comp_lookup.get(rev.company_id)
        b_id = comp.original_id if comp else ""
        b_name = comp.original_name if comp else ""
        conf_vals = " | ".join(json.loads(rev.conflicting_values or "[]"))
        urls = " | ".join(json.loads(rev.source_urls or "[]"))

        ws_review.append([
            b_id, b_name, rev.field_name, conf_vals, urls,
            rev.reason, rev.suggested_action, rev.review_status
        ])
        curr_row = ws_review.max_row
        for c_idx in range(1, len(headers_sheet3) + 1):
            cell = ws_review.cell(row=curr_row, column=c_idx)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER

    style_header_row(ws_review, len(headers_sheet3))
    auto_fit_columns(ws_review)

    # ==========================================
    # SHEET 4: Research Summary
    # ==========================================
    ws_summary = wb.create_sheet(title="Research Summary")
    ws_summary.append(["AI LEAD INTELLIGENCE - RESEARCH SUMMARY REPORT"])
    ws_summary.append(["Generated by Tex Web Solutions Pvt. Ltd.", f"Date: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}"])
    ws_summary.append([])
    
    summary_data = [
        ["Job ID", job.id],
        ["Filename", job.filename],
        ["Total Businesses Uploaded", job.total_companies],
        ["Total Businesses Processed", job.processed_companies],
        ["Successfully Enriched", job.enriched_count],
        ["Fully Verified", job.fully_verified_count],
        ["Partially Verified", job.partially_verified_count],
        ["Unverified", job.unverified_count],
        ["Not Found / Empty", job.not_found_count],
        ["Failed / Errored", job.failed_count],
        ["Requires Manual Review", job.manual_review_count],
        ["Job Status", job.status.upper()],
        ["Job Started", job.created_at.strftime("%Y-%m-%d %H:%M:%S") if job.created_at else ""],
        ["Job Last Updated", job.updated_at.strftime("%Y-%m-%d %H:%M:%S") if job.updated_at else ""],
    ]
    for row in summary_data:
        ws_summary.append(row)

    # Title styling
    ws_summary.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    ws_summary.cell(row=2, column=1).font = Font(name="Calibri", size=10, italic=True, color="64748B")
    for r in range(4, 4 + len(summary_data)):
        ws_summary.cell(row=r, column=1).font = Font(name="Calibri", size=11, bold=True)
        ws_summary.cell(row=r, column=2).font = Font(name="Calibri", size=11)
        ws_summary.cell(row=r, column=1).border = THIN_BORDER
        ws_summary.cell(row=r, column=2).border = THIN_BORDER

    auto_fit_columns(ws_summary)

    # Save to export directory
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    clean_fn = Path(job.filename).stem.replace(" ", "_")
    export_filename = f"Enriched_{clean_fn}_{timestamp}.xlsx"
    export_path = EXPORT_DIR / export_filename
    wb.save(export_path)

    return export_path

def export_job_to_csv(job: Job, companies: list[Company]) -> Path:
    """
    Exports Sheet 1 data directly to CSV.
    """
    data = []
    for comp in companies:
        add_emails = ", ".join(json.loads(comp.additional_business_emails or "[]"))
        add_phones = ", ".join(json.loads(comp.additional_business_phones or "[]"))
        missing = ", ".join(json.loads(comp.missing_fields or "[]"))
        last_chk = comp.last_checked.strftime("%Y-%m-%d %H:%M:%S") if comp.last_checked else ""

        row = {
            "id": comp.original_id,
            "name": comp.original_name,
            "trades": comp.original_trades,
            "city": comp.original_city,
            "state": comp.original_state,
            "phone": comp.original_phone,
            "webssite url": comp.original_website,
            "claimed": comp.original_claimed,
            "Official Business Name": comp.official_name or comp.original_name,
            "Business Category": comp.business_category or comp.trades,
            "Business Description": comp.business_description,
            "Full Business Address": comp.address,
            "ZIP Code": comp.zip_code,
            "Country": comp.country,
            "Year Established": comp.year_established,
            "Company Size": comp.company_size,
            "Official Website": comp.official_website,
            "Website Status": comp.website_status,
            "Tech Stack / CMS": comp.tech_stack_cms or "Custom / Unknown",
            "SSL Secured": "Yes" if comp.has_ssl else "No",
            "Mobile Responsive": "Yes" if comp.is_mobile_responsive else "No",
            "Has Analytics": "Yes" if comp.has_analytics else "No",
            "SEO Health Score": comp.seo_audit_score or 0,
            "Redesign Opportunity": comp.redesign_opportunity or "Medium",
            "Owner Name": comp.owner_name,
            "Owner Job Title": comp.owner_title,
            "CEO Name": comp.ceo_name,
            "CEO Job Title": comp.ceo_title,
            "Founder Name": comp.founder_name,
            "Founder Job Title": comp.founder_title,
            "Co-Founder Name": comp.cofounder_name,
            "Co-Founder Job Title": comp.cofounder_title,
            "Primary Decision Maker": comp.primary_decision_maker,
            "Primary Decision Maker Title": comp.primary_decision_maker_title,
            "Business Email": comp.business_email,
            "Additional Business Emails": add_emails,
            "Business Phone": comp.business_phone,
            "Additional Business Phone": add_phones,
            "Personal Professional Email": comp.personal_email,
            "Public Professional Phone": comp.personal_phone,
            "Phone Line Type": comp.phone_line_type or "Landline / Office",
            "Email Verification Status": comp.email_verification_status,
            "Mailbox Provider": comp.email_mailbox_provider or "Custom",
            "Personal LinkedIn": comp.personal_linkedin,
            "Business LinkedIn": comp.business_linkedin,
            "Business Instagram": comp.business_instagram,
            "Business Facebook": comp.business_facebook,
            "Owner Verification Status": comp.owner_status,
            "CEO Verification Status": comp.ceo_status,
            "Founder Verification Status": comp.founder_status,
            "Co-Founder Verification Status": comp.cofounder_status,
            "Business Email Status": comp.business_email_status,
            "Business Phone Status": comp.business_phone_status,
            "Personal Email Status": comp.personal_email_status,
            "Personal Phone Status": comp.personal_phone_status,
            "Personal LinkedIn Status": comp.personal_linkedin_status,
            "Business LinkedIn Status": comp.business_linkedin_status,
            "Instagram Status": comp.business_instagram_status,
            "Facebook Status": comp.business_facebook_status,
            "Overall Confidence Score": comp.overall_confidence_score,
            "Overall Verification Status": comp.overall_verification_status,
            "Missing Fields": missing,
            "Last Checked": last_chk,
            "Research Notes": comp.research_notes,
            "Personalized Outreach Pitch": comp.ai_personalized_pitch or ""
        }
        data.append(row)

    df = pd.DataFrame(data)
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    clean_fn = Path(job.filename).stem.replace(" ", "_")
    export_filename = f"Enriched_{clean_fn}_{timestamp}.csv"
    export_path = EXPORT_DIR / export_filename
    df.to_csv(export_path, index=False, encoding="utf-8-sig")
    return export_path



def export_automation_to_excel(auto_job: AutomationJob, companies: list[Company]) -> Path:
    """
    Exports Automated Lead Generation Run into a professional 6-Worksheet Excel Workbook:
    1. All Companies
    2. Decision Makers
    3. Verified Emails
    4. Missing Websites
    5. Failed Enrichment
    6. Automation Summary
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove blank default

    # Standard Lead Columns
    headers = [
        "Lead ID", "Company Name", "Industry / Trade", "Opportunity Score", "Score Reason",
        "City", "State", "Address", "ZIP Code", "Latitude", "Longitude", "Distance (Miles)",
        "Official Website", "Website Status", "Tech Stack / CMS", "SSL Secured", "Mobile Responsive",
        "SEO Score", "Redesign Opportunity", "Primary Decision Maker", "Title",
        "Owner Name", "CEO Name", "Founder Name",
        "Business Email", "Email Status", "Mailbox Provider",
        "Business Phone", "Phone Line Type", "Personal Email", "Personal Phone",
        "Personal LinkedIn", "Business LinkedIn", "Instagram", "Facebook",
        "Discovery Provider", "Directory Profile URL", "Directory Description", "Verification Status", "Personalized Outreach Pitch"
    ]

    def create_lead_row(c: Company) -> list[Any]:
        lead_id = c.original_id or (c.id[:8] if c.id else "LEAD")
        return [
            lead_id,
            c.official_name or c.original_name,
            c.trades or c.original_trades,
            c.lead_score or 0,
            c.lead_score_reason or "N/A",
            c.city or c.original_city or "N/A",
            c.state or c.original_state or "N/A",
            c.address or "N/A",
            c.zip_code or "N/A",
            c.latitude if c.latitude is not None else "N/A",
            c.longitude if c.longitude is not None else "N/A",
            c.distance_from_center_miles or 0.0,
            c.official_website or c.original_website or "N/A",
            c.website_status or "N/A",
            c.tech_stack_cms or "Custom",
            "Yes" if c.has_ssl else "No",
            "Yes" if c.is_mobile_responsive else "No",
            c.seo_audit_score or 0,
            c.redesign_opportunity or "Medium",
            c.primary_decision_maker or "N/A",
            c.primary_decision_maker_title or "N/A",
            c.owner_name or "N/A",
            c.ceo_name or "N/A",
            c.founder_name or "N/A",
            c.business_email or "N/A",
            c.email_verification_status or "Unknown",
            c.email_mailbox_provider or "Custom",
            c.business_phone or "N/A",
            c.phone_line_type or "Landline / Office",
            c.personal_email or "N/A",
            c.personal_phone or "N/A",
            c.personal_linkedin or "N/A",
            c.business_linkedin or "N/A",
            c.business_instagram or "N/A",
            c.business_facebook or "N/A",
            c.discovery_provider or "Automated",
            c.directory_profile_url or "N/A",
            c.directory_description or "N/A",
            c.overall_verification_status or "Unverified",
            c.ai_personalized_pitch or "N/A"
        ]

    def populate_sheet(ws, row_list: list[list[Any]]):
        ws.append(headers)
        for r in row_list:
            ws.append(r)
            curr = ws.max_row
            for c_idx in range(1, len(r) + 1):
                cell = ws.cell(row=curr, column=c_idx)
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                v = str(cell.value or "")
                if v.startswith("http://") or v.startswith("https://"):
                    cell.hyperlink = v
                    cell.font = LINK_FONT
        style_header_row(ws, len(headers))
        auto_fit_columns(ws)

    # 1. All Companies
    ws_all = wb.create_sheet(title="All Companies")
    all_rows = [create_lead_row(c) for c in companies]
    populate_sheet(ws_all, all_rows)

    # 2. Decision Makers
    ws_dm = wb.create_sheet(title="Decision Makers")
    dm_rows = [create_lead_row(c) for c in companies if (c.owner_name or c.ceo_name or c.founder_name or c.primary_decision_maker)]
    populate_sheet(ws_dm, dm_rows)

    # 3. Verified Emails
    ws_emails = wb.create_sheet(title="Verified Emails")
    email_rows = [create_lead_row(c) for c in companies if c.email_verification_status == "Valid"]
    populate_sheet(ws_emails, email_rows)

    # 4. Missing Websites
    ws_missing = wb.create_sheet(title="Missing Websites")
    missing_rows = [create_lead_row(c) for c in companies if c.website_status in ["No Official Website", "Broken", "N/A", "Unverified"]]
    populate_sheet(ws_missing, missing_rows)

    # 5. Failed Enrichment
    ws_failed = wb.create_sheet(title="Failed Enrichment")
    failed_rows = [create_lead_row(c) for c in companies if c.overall_verification_status == "Failed"]
    populate_sheet(ws_failed, failed_rows)

    # 6. Automation Summary
    ws_summary = wb.create_sheet(title="Automation Summary")
    ws_summary.append(["AI LEAD INTELLIGENCE — AUTOMATED LEAD GENERATION REPORT"])
    ws_summary.append(["Tex Web Solutions Pvt. Ltd.", f"Run Date: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}"])
    ws_summary.append([])

    cfg = json.loads(auto_job.search_config_json or "{}")
    summary_data = [
        ["Automation Name", auto_job.name],
        ["Status", auto_job.status.upper()],
        ["Target States", ", ".join(cfg.get("states", []))],
        ["Search Radius", f"{cfg.get('radius_miles', 50)} Miles"],
        ["Target Industries", ", ".join(cfg.get("industries", []))],
        ["Total Target Requested", auto_job.total_target],
        ["Total Businesses Discovered", auto_job.total_discovered],
        ["Unique Businesses Found", auto_job.unique_discovered],
        ["Duplicate Records Filtered", auto_job.duplicate_count],
        ["Successfully Enriched", auto_job.total_enriched],
        ["Verified MX Emails Found", auto_job.verified_emails_count],
        ["Failed / Errored Records", auto_job.failed_count],
        ["Run Started", auto_job.start_time.strftime("%Y-%m-%d %H:%M:%S") if auto_job.start_time else "N/A"],
        ["Run Completed", auto_job.end_time.strftime("%Y-%m-%d %H:%M:%S") if auto_job.end_time else "In Progress"]
    ]
    for row in summary_data:
        ws_summary.append(row)

    ws_summary.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    ws_summary.cell(row=2, column=1).font = Font(name="Calibri", size=10, italic=True, color="64748B")
    for r in range(4, 4 + len(summary_data)):
        ws_summary.cell(row=r, column=1).font = Font(name="Calibri", size=11, bold=True)
        ws_summary.cell(row=r, column=2).font = Font(name="Calibri", size=11)
        ws_summary.cell(row=r, column=1).border = THIN_BORDER
        ws_summary.cell(row=r, column=2).border = THIN_BORDER
    auto_fit_columns(ws_summary)

    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    clean_fn = auto_job.name.replace(" ", "_")
    export_filename = f"Automation_{clean_fn}_{timestamp}.xlsx"
    export_path = EXPORT_DIR / export_filename
    wb.save(export_path)
    return export_path


def export_automation_to_csv(auto_job: AutomationJob, companies: list[Company]) -> Path:
    """Exports automation discovered leads directly to CSV."""
    data = []
    for c in companies:
        lead_id = c.original_id or (c.id[:8] if c.id else "LEAD")
        data.append({
            "Lead ID": lead_id,
            "Company Name": c.official_name or c.original_name,
            "Industry": c.trades or c.original_trades,
            "Opportunity Score": c.lead_score or 0,
            "Score Reason": c.lead_score_reason or "N/A",
            "City": c.city or c.original_city or "N/A",
            "State": c.state or c.original_state or "N/A",
            "Address": c.address or "N/A",
            "ZIP Code": c.zip_code or "N/A",
            "Latitude": c.latitude if c.latitude is not None else "N/A",
            "Longitude": c.longitude if c.longitude is not None else "N/A",
            "Distance (Miles)": c.distance_from_center_miles or 0.0,
            "Website": c.official_website or c.original_website or "N/A",
            "Website Status": c.website_status or "N/A",
            "Tech Stack / CMS": c.tech_stack_cms or "Custom",
            "SSL Secured": "Yes" if c.has_ssl else "No",
            "Mobile Responsive": "Yes" if c.is_mobile_responsive else "No",
            "SEO Score": c.seo_audit_score or 0,
            "Redesign Opportunity": c.redesign_opportunity or "Medium",
            "Decision Maker": c.primary_decision_maker or c.owner_name or "N/A",
            "Title": c.primary_decision_maker_title or "N/A",
            "Business Email": c.business_email or "N/A",
            "Email Status": c.email_verification_status or "Unknown",
            "Mailbox Provider": c.email_mailbox_provider or "Custom",
            "Business Phone": c.business_phone or "N/A",
            "Phone Line Type": c.phone_line_type or "Landline / Office",
            "Personal Email": c.personal_email or "N/A",
            "Personal Phone": c.personal_phone or "N/A",
            "Personal LinkedIn": c.personal_linkedin or "N/A",
            "Business LinkedIn": c.business_linkedin or "N/A",
            "Instagram": c.business_instagram or "N/A",
            "Facebook": c.business_facebook or "N/A",
            "Provider": c.discovery_provider or "Automated",
            "Directory Profile URL": c.directory_profile_url or "N/A",
            "Directory Description": c.directory_description or "N/A",
            "Verification Status": c.overall_verification_status or "Unverified",
            "Personalized Pitch": c.ai_personalized_pitch or "N/A"
        })

    df = pd.DataFrame(data)
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    clean_fn = auto_job.name.replace(" ", "_")
    export_filename = f"Automation_{clean_fn}_{timestamp}.csv"
    export_path = EXPORT_DIR / export_filename
    df.to_csv(export_path, index=False, encoding="utf-8-sig")
    return export_path


def export_national_roofing_job_to_excel(job: NationalRoofingJob, companies: list[Company]) -> Path:
    """
    Generates an 8-sheet OpenPyXL Excel workbook for National Roofing Directory scraping results.
    Sheets:
    1. {State} Roofing Leads
    2. City Summary
    3. Scrape Summary
    4. Gmail Found
    5. Emails Found
    6. Websites Found
    7. Missing Websites
    8. Failed Profiles
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    state_name = (job.state or "Texas").strip()
    main_sheet_title = f"{state_name} Roofing Leads"[:31]

    # Common Table Headers
    lead_headers = [
        "Company Name", "Owner / Contact Name", "Trades / Services", "City", "State",
        "ZIP Code", "Business Address", "Phone", "Email", "Gmail Address",
        "Company Website URL", "Website Status", "Profile URL", "Verification Status"
    ]

    def build_lead_row(c: Company) -> list[Any]:
        return [
            c.official_name or c.original_name or "N/A",
            c.owner_name or "N/A",
            c.trades or "Roofing Contractors",
            c.city or "N/A",
            c.state or "TX",
            c.zip_code or "N/A",
            c.address or "N/A",
            c.business_phone or "N/A",
            c.business_email or "N/A",
            c.gmail_address or "N/A",
            c.official_website or "N/A",
            c.website_status or "No Website Found",
            c.directory_profile_url or "N/A",
            c.overall_verification_status or "Unverified"
        ]

    # ==========================================
    # SHEET 1: {State} Roofing Leads
    # ==========================================
    ws_main = wb.create_sheet(title=main_sheet_title)
    ws_main.append(lead_headers)
    for c in companies:
        ws_main.append(build_lead_row(c))
    style_header_row(ws_main, len(lead_headers))
    auto_fit_columns(ws_main)

    # ==========================================
    # SHEET 2: City Summary
    # ==========================================
    ws_city = wb.create_sheet(title="City Summary")
    city_headers = ["City", "Total Contractors", "Emails Found", "Gmail Found", "Websites Found", "Missing Websites"]
    ws_city.append(city_headers)

    city_counts: dict[str, dict[str, int]] = {}
    for c in companies:
        c_name = c.city or "Unknown"
        if c_name not in city_counts:
            city_counts[c_name] = {"total": 0, "emails": 0, "gmail": 0, "websites": 0, "missing_web": 0}
        city_counts[c_name]["total"] += 1
        if c.business_email and c.business_email != "N/A":
            city_counts[c_name]["emails"] += 1
        if c.gmail_address and c.gmail_address != "N/A":
            city_counts[c_name]["gmail"] += 1
        if c.official_website and c.official_website != "N/A":
            city_counts[c_name]["websites"] += 1
        else:
            city_counts[c_name]["missing_web"] += 1

    for c_name, stats in sorted(city_counts.items()):
        ws_city.append([c_name, stats["total"], stats["emails"], stats["gmail"], stats["websites"], stats["missing_web"]])
    style_header_row(ws_city, len(city_headers))
    auto_fit_columns(ws_city)

    # ==========================================
    # SHEET 3: Scrape Summary
    # ==========================================
    ws_summary = wb.create_sheet(title="Scrape Summary")
    ws_summary.append(["National Roofing Directory Scrape Summary", "Value"])

    missing_web_count = sum(1 for c in companies if not c.official_website or c.official_website == "N/A")

    summary_rows = [
        ("Directory Source", "National Roofing Directory (https://www.nationalroofingdirectory.com/)"),
        ("Selected State", state_name),
        ("City Pages Discovered", job.city_pages_discovered),
        ("City Pages Completed", job.city_pages_completed),
        ("Unique Profile URLs Discovered", job.unique_profiles_count or len(companies)),
        ("Profiles Successfully Processed", job.profiles_processed or len(companies)),
        ("Profiles with Phone", job.phones_found or sum(1 for c in companies if c.business_phone and c.business_phone != "N/A")),
        ("Profiles with Public Email", job.emails_found or sum(1 for c in companies if c.business_email and c.business_email != "N/A")),
        ("Profiles with Gmail Addresses", job.gmail_found or sum(1 for c in companies if c.gmail_address and c.gmail_address != "N/A")),
        ("Profiles with Company Websites", job.websites_found or sum(1 for c in companies if c.official_website and c.official_website != "N/A")),
        ("Profiles without Websites", missing_web_count),
        ("Failed Profile URLs", job.failed_profiles_count),
        ("Execution Status", job.status.upper()),
        ("Start Time", str(job.start_time) if job.start_time else "N/A"),
        ("End Time", str(job.end_time) if job.end_time else "In Progress"),
        ("Export Generated At", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"))
    ]
    for k, v in summary_rows:
        ws_summary.append([k, v])
    style_header_row(ws_summary, 2)
    auto_fit_columns(ws_summary, min_width=25, max_width=65)

    # ==========================================
    # SHEET 4: Gmail Found
    # ==========================================
    ws_gmail = wb.create_sheet(title="Gmail Found")
    ws_gmail.append(lead_headers)
    for c in companies:
        if c.gmail_address and c.gmail_address != "N/A":
            ws_gmail.append(build_lead_row(c))
    style_header_row(ws_gmail, len(lead_headers))
    auto_fit_columns(ws_gmail)

    # ==========================================
    # SHEET 5: Emails Found
    # ==========================================
    ws_emails = wb.create_sheet(title="Emails Found")
    ws_emails.append(lead_headers)
    for c in companies:
        has_biz = bool(c.business_email and c.business_email != "N/A")
        has_gm = bool(c.gmail_address and c.gmail_address != "N/A")
        if has_biz or has_gm:
            ws_emails.append(build_lead_row(c))
    style_header_row(ws_emails, len(lead_headers))
    auto_fit_columns(ws_emails)

    # ==========================================
    # SHEET 6: Websites Found
    # ==========================================
    ws_web = wb.create_sheet(title="Websites Found")
    ws_web.append(lead_headers)
    for c in companies:
        if c.official_website and c.official_website != "N/A":
            ws_web.append(build_lead_row(c))
    style_header_row(ws_web, len(lead_headers))
    auto_fit_columns(ws_web)

    # ==========================================
    # SHEET 7: Missing Websites
    # ==========================================
    ws_missing = wb.create_sheet(title="Missing Websites")
    ws_missing.append(lead_headers)
    for c in companies:
        if not c.official_website or c.official_website == "N/A":
            ws_missing.append(build_lead_row(c))
    style_header_row(ws_missing, len(lead_headers))
    auto_fit_columns(ws_missing)

    # ==========================================
    # SHEET 8: Failed Profiles
    # ==========================================
    ws_failed = wb.create_sheet(title="Failed Profiles")
    ws_failed.append(["Profile URL", "Failure Reason", "Timestamp"])
    
    # Load failed profiles from checkpoint_json
    try:
        cp = json.loads(job.checkpoint_json or "{}")
        failed_list = cp.get("failed_urls", [])
        for f_url in failed_list:
            ws_failed.append([f_url, "HTTP 404 / Parse Error", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")])
    except Exception:
        pass
    style_header_row(ws_failed, 3)
    auto_fit_columns(ws_failed)

    # Save to file
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    clean_state = state_name.replace(" ", "_")
    export_filename = f"National_Roofing_{clean_state}_{timestamp}.xlsx"
    export_path = EXPORT_DIR / export_filename
    wb.save(export_path)
    return export_path


def export_national_roofing_job_to_csv(job: NationalRoofingJob, companies: list[Company]) -> Path:
    """Exports National Roofing Directory results to CSV."""
    data = []
    for c in companies:
        data.append({
            "Company Name": c.official_name or c.original_name or "N/A",
            "Owner / Contact Name": c.owner_name or "N/A",
            "Trades / Services": c.trades or "Roofing Contractors",
            "City": c.city or "N/A",
            "State": c.state or "TX",
            "ZIP Code": c.zip_code or "N/A",
            "Business Address": c.address or "N/A",
            "Phone": c.business_phone or "N/A",
            "Email": c.business_email or "N/A",
            "Gmail Address": c.gmail_address or "N/A",
            "Company Website URL": c.official_website or "N/A",
            "Website Status": c.website_status or "No Website Found",
            "Profile URL": c.directory_profile_url or "N/A",
            "Verification Status": c.overall_verification_status or "Unverified"
        })

    df = pd.DataFrame(data)
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    clean_state = (job.state or "Texas").replace(" ", "_")
    export_filename = f"National_Roofing_{clean_state}_{timestamp}.csv"
    export_path = EXPORT_DIR / export_filename
    df.to_csv(export_path, index=False, encoding="utf-8-sig")
    return export_path



